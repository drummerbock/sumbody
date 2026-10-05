"""Import the original SUMbody / My Tip Jar workbook.
Usage: python scripts/import_legacy.py /path/to/workbook.xlsx user@example.com
The workbook itself should never be committed to this public repository.
"""
import sys
from pathlib import Path
from datetime import datetime, date
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from app import app, db, User, LifeArea, SumEntry

def as_date(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    if isinstance(v, (int,float)):
        return datetime(1899,12,30).date() + __import__("datetime").timedelta(days=int(v))
    if isinstance(v,str):
        for fmt in ("%m/%d/%Y","%Y-%m-%d"):
            try:return datetime.strptime(v.strip(),fmt).date()
            except ValueError:pass
    return None

def main(path,email):
    wb=load_workbook(path,read_only=True,data_only=True)
    ws=wb["Form Responses 1"]
    with app.app_context():
        user=User.query.filter_by(email=email.lower()).first()
        if not user: raise SystemExit("Create the user first with: flask --app app create-user")
        areas={a.name:a for a in LifeArea.query.filter_by(user_id=user.id).all()}
        added=skipped=0
        for row in ws.iter_rows(min_row=2,values_only=True):
            timestamp,description,area_name,day,value,*rest=row
            if not description or not area_name or not day or value is None: skipped+=1; continue
            d=as_date(day)
            try: score=int(value)
            except (ValueError,TypeError): skipped+=1; continue
            # Preserve standard 1–5 SUMs in v1; exceptional experimental values are skipped for review.
            if not d or score not in range(1,6): skipped+=1; continue
            name=str(area_name).strip()
            if name not in areas:
                areas[name]=LifeArea(user_id=user.id,name=name,active=False); db.session.add(areas[name]); db.session.flush()
            comm=rest[0] if rest else None
            comm_bool=True if str(comm).strip().lower()=="yes" else False if str(comm).strip().lower()=="no" else None
            db.session.add(SumEntry(user_id=user.id,life_area_id=areas[name].id,description=str(description).strip(),value=score,communicated=comm_bool,occurred_on=d,legacy_timestamp=str(timestamp) if timestamp else None))
            added+=1
        db.session.commit()
        print(f"Imported {added} SUMs; skipped {skipped} rows for review.")

if __name__=="__main__":
    if len(sys.argv)!=3: raise SystemExit("Usage: python scripts/import_legacy.py workbook.xlsx user@example.com")
    main(sys.argv[1],sys.argv[2])
